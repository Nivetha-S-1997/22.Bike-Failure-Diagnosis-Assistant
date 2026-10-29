#Importing the Libraries:
from helper_func import chat_with_model

from docx import Document as DocxDocument

from langchain_google_genai import GoogleGenerativeAIEmbeddings
from langchain_chroma import Chroma
from langchain_core.documents import Document

#Document & Embeddings Loading:
doc=DocxDocument('Common bike problems- RAG.docx') #Reading word document

para_data=[] #Document reads plain texts (paragraphs)
table_data=[] #Document reads tables also

for para in doc.paragraphs: #Loops through the texts other than table
    if para.text.strip():
        para_data.append(para.text)

for table in doc.tables: #Loops through each table
    for row in table.rows: #Loops through each row in the table
        row_text=[cell.text.strip() for cell in row.cells] #Loops through each cell in the table & stores the texts of each cell
        table_data.append('|'.join(row_text)) #

documents=[Document(page_content=text) for text in (para_data + table_data)]
print(f"Total documents for embedding: {len(documents)}", flush=True)

retriever = None

def get_retriever():
    global retriever

    if retriever is None:
        embeddings = GoogleGenerativeAIEmbeddings(model="models/gemini-embedding-001", batch_size=20)

        db = Chroma(collection_name="bike_problems", embedding_function=embeddings, persist_directory="./chroma_db")

        stored_count = db._collection.count()
        expected_count = len(documents)

        print(f"Stored documents: {stored_count}", flush=True)

        if stored_count != expected_count:
            raise RuntimeError(
                f"ChromaDB contains {stored_count} documents; "
                f"expected {expected_count}. "
                "The saved vector database is missing or incomplete."
            )

        print("Using saved ChromaDB embeddings.", flush=True)
        
        print(f"Documents stored in ChromaDB: {db._collection.count()}", flush=True)

        retriever = db.as_retriever()

    return retriever

chat_sessions={}

#Main function:
def process_user_message(user_text, session_id):
    chat_history = chat_sessions.setdefault(session_id, [])
    docs = get_retriever().invoke(user_text)
    context = "\n".join([d.page_content for d in docs])

    history_text = "\n".join(
    [f"User: {item['user']}\nAssistant: {item['bot']}"
     for item in chat_history]
)
    # Create chatbot response:
    System_instruction=System_instruction = '''You are a helpful assistant that provides solutions to customers' 
    bike faults, similar to a mechanic.

    Use the retrieved context and previous conversation to understand the user's problem.
    Respond naturally and do not mention the retrieved context or knowledge base to the user.

    Follow-up questions:
    - Ask follow-up questions when important information is missing.
    - Ask only the questions necessary to understand the problem.
    - Avoid asking unnecessary or repetitive questions.
    - Do not repeat questions that have already been answered.
    - Once sufficient information is available, provide a possible diagnosis and recommended actions.
    - If the exact cause remains uncertain, explain the uncertainty and recommend professional inspection when appropriate.
    - If the user provides additional symptoms, reassess the diagnosis based on the new information.

    Response format:
    When providing a diagnosis, organize the response under these headings:

    **Possible Diagnosis**
    Explain the likely causes based on the symptoms and available information.
    Do not present uncertain causes as confirmed faults.

    **Recommended Actions**
    Provide practical, safe troubleshooting steps that the user can follow.
    Recommend a qualified mechanic when professional inspection or repair is necessary.

    When asking follow-up questions, respond naturally without forcing the diagnosis format.

    Safety instructions:
    - If the user's message requests harmful, violent, sexual, or hateful
    content unrelated to bike troubleshooting, politely refuse.
    - Do not provide dangerous repair instructions involving live electrical
    connections, fuel systems, or other hazardous procedures.
    - Recommend a qualified mechanic when a repair could be unsafe.'''

    message = [
    {
        'role': 'user',
        'parts': [{
            'text': System_instruction +
                    "\n\nPrevious conversation:\n" + history_text +
                    "\n\nCurrent user message:\n" + user_text
        }]
    }
]

    bot_reply=chat_with_model(message, context)

    chat_history.append({
    "user": user_text,
    "bot": bot_reply
})

    return bot_reply